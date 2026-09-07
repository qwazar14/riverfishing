# -*- coding: utf-8 -*-
"""§species-table (0.10): the species live in a spreadsheet now. This reads the author's table and writes
everything a species is made of into every tree: profile JSON, names and descriptions in three languages,
the item model (the pike's, re-pointed), the roster in ModItems, the fishes tag, the hybrid parents.

    py tools/import_species_xlsx.py <table.xlsx> [tree ...]      (default: the three trees)

Rules the table carries and the profile no longer does: distance_pref, ideal.rod / rig / reel_size /
reel_tolerance and hook.tolerance are DROPPED from every profile (they left the bite engine in the same
change). min_angler_level is recomputed on a 0-50 scale: the hand-set 0-12 of the old species stretched,
the new species placed by their three nearest old neighbours in (log weight, fight strength).
Fields the table does not carry (legendary, breeds_with, gynogenesis, bed, display) are preserved."""
import io, json, math, os, re, shutil, sys
import openpyxl

MAIN = r"C:/Users/Qwazar/VS Code Projects/fishing mod"
TREES = [MAIN, r"C:/Users/Qwazar/wt/rf1201", r"C:/Users/Qwazar/wt/rf26"]
A = "common/src/main/resources/assets/riverfishing"
D = "common/src/main/resources/data/riverfishing"
TAG = "§species-table (0.10)"
DONOR = "pike"
OLD_LEVEL_MAX = 12.0   # the hand-set ladder topped out here; the new one runs to 50

# the crosses the table's hybrids come from — the parents get breeds_with, the hybrid gets hybrid_of
HYBRIDS = {
    "cutbow": ("cutthroat_trout", "rainbow_trout"), "splake": ("brook_trout", "lake_trout"),
    "tiger_muskie": ("muskellunge", "northern_pike"), "saugeye": ("sauger", "walleye"),
    "wiper": ("white_bass", "striped_bass"), "bester_sturgeon": ("beluga", "sterlet"),
    "black_white_crappie_hybrid": ("black_crappie", "white_crappie"),
    "channel_blue_catfish_hybrid": ("channel_catfish", "blue_catfish"),
    "largemouth_smallmouth_hybrid": ("largemouth_bass", "smallmouth_bass"),
    "smallmouth_spotted_bass_hybrid": ("smallmouth_bass", "spotted_bass"),
    "bluegill_redear_hybrid": ("bluegill", "redear_sunfish"), "bluegill_redbreast_hybrid": ("bluegill", "redbreast_sunfish"),
    "bluegill_pumpkinseed_hybrid": ("bluegill", "pumpkinseed"), "bluegill_green_sunfish_hybrid": ("bluegill", "green_sunfish"),
    "redear_green_sunfish_hybrid": ("redear_sunfish", "green_sunfish"),
    "chinook_coho_hybrid": ("chinook_salmon", "coho_salmon"), "chinook_pink_hybrid": ("chinook_salmon", "pink_salmon"),
    "brook_bull_trout_hybrid": ("brook_trout", "bull_trout"), "dolly_bull_trout_hybrid": ("dolly_varden", "bull_trout"),
    "common_carp_gibel_hybrid": ("carp", "crucian_carp"), "kaluga_sterlet_hybrid": ("kaluga_sturgeon", "sterlet"),
    "kuria_labeo_catla_hybrid": ("kuria_labeo", "catla"), "rohu_catla_hybrid": ("rohu", "catla"),
    "rohu_kuria_labeo_hybrid": ("rohu", "kuria_labeo"), "nile_blue_tilapia_hybrid": ("nile_tilapia", "blue_tilapia"),
    "nile_mozambique_tilapia_hybrid": ("nile_tilapia", "mozambique_tilapia"),
    "blue_mozambique_tilapia_hybrid": ("blue_tilapia", "mozambique_tilapia"),
    "longnose_alligator_gar_hybrid": ("longnose_gar", "alligator_gar"),
    "bream_roach_hybrid": ("bream", "roach"), "roach_rudd_hybrid": ("roach", "rudd"),
}
GROUP_NAMES = {
    "panfish": ("Panfish", "Солнечные окуни", "Сонячні окуні"), "catfish": ("Catfishes", "Сомы", "Соми"),
    "cichlid": ("Cichlids", "Цихлиды", "Цихліди"), "characin": ("Characins", "Харациновые", "Харацинові"),
    "exotic": ("Exotics", "Экзотика", "Екзотика"), "ray": ("Rays", "Скаты", "Скати"),
}
DIET_NAMES = {
    "predator": ("Predator", "Хищник", "Хижак"), "omnivore": ("Omnivore", "Всеядная", "Всеїдна"),
    "peaceful": ("Peaceful", "Мирная", "Мирна"), "insectivore": ("Insect feeder", "Насекомоядная", "Комахоїдна"),
}
PROVINCE_NAMES = {"afrotropical": ("Afrotropical", "Афротропика", "Афротропіка")}
LOCS = ("en_us", "ru_ru", "uk_ua")


def rd(p): return io.open(p, encoding="utf-8").read()
def wr(p, s):
    os.makedirs(os.path.dirname(p), exist_ok=True)
    io.open(p, "w", encoding="utf-8", newline="\n").write(s)
def jload(p): return json.load(io.open(p, encoding="utf-8"))
def jdump(p, d): wr(p, json.dumps(d, ensure_ascii=False, indent=2) + "\n")


def weights(cell):
    """'river 1.3; lake 0.9' -> {'river': 1.3, 'lake': 0.9}"""
    out = {}
    if cell is None: return out
    for part in str(cell).replace(",", ";").split(";"):
        part = part.strip()
        if not part: continue
        bits = part.split()
        if len(bits) != 2: raise ValueError("bad weight pair %r" % part)
        out[bits[0]] = float(bits[1])
    return out


def listing(cell):
    return [x.strip() for x in str(cell).replace(";", ",").split(",") if x.strip()] if cell else []


def num(v, cast=float):
    if v is None or v == "": return None
    return cast(v)


def read_table(path):
    ws = openpyxl.load_workbook(path, data_only=True)["Species"]
    hdr = [c.value for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=2):
        row = dict(zip(hdr, [c.value for c in r]))
        sid = row.get("ID")
        if not sid or str(sid).strip() in ("snake_case, unique", "ID"): continue
        row["ID"] = str(sid).strip()
        rows.append(row)
    ids = [r["ID"] for r in rows]
    assert len(ids) == len(set(ids)), "duplicate ids"
    return rows


def build_profile(row, old):
    """The profile for one row, on top of what the tree already had for it (or the donor's shape)."""
    p = dict(old) if old else {}
    sid = row["ID"]
    p["display"] = old.get("display") if old and old.get("display") else row["English name"]
    p["latin"] = str(row["Scientific name"] or "").strip()
    p["group"] = str(row["Group"]).strip()
    p["diet"] = str(row["Diet profile"]).strip()
    p["spawn"] = {"season": str(row["Spawn season"]).strip()}
    if row.get("Spawn sub"): p["spawn"]["sub"] = str(row["Spawn sub"]).strip()
    p["water_bodies"] = weights(row["Water bodies"])
    maxg = int(round(float(row["max. weight (kg)"]) * 1000))
    p["weight_g"] = {"min": int(round(num(row["Min weight (g)"]))), "max": maxg, "mean": int(round(num(row["Mean weight (g)"])))}
    p["length_cm"] = {"min": num(row["Min length (cm)"]), "max": num(row["max. length (cm)"])}
    p["fight"] = {"strength": num(row["Strength"]), "stamina": num(row["Stamina"]), "runs": int(num(row["Runs"], int)),
                  "pattern": str(row["Behaviour"]).strip(), "aggression": num(row["Aggression"])}
    ideal = dict(old.get("ideal", {})) if old else {}
    for k in ("rod", "rig", "reel_size", "reel_tolerance"): ideal.pop(k, None)
    line = dict(ideal.get("line", {}))
    line["type"] = str(row["Line type"]).strip()
    line["diameter_mm"] = num(row["Line diameter (mm)"])
    line.setdefault("tolerance_mm", 0.06)
    ideal["line"] = line
    ideal["groundbait"] = {"fraction": num(row["Groundbait fraction"]), "nutrition": num(row["Groundbait nutrition"])}
    ideal["bait"] = weights(row["Baits & Lures"])
    ideal["hook"] = {"ideal": int(num(row["Hook ideal"], int))}
    ideal["requires_leader"] = str(row["Requires leader"]).strip().lower() in ("yes", "да", "так", "true", "1")
    p["ideal"] = ideal
    p["season"] = weights(row["Season"])
    p["time"] = weights(row["Time"])
    p["weather"] = weights(row["Weather"])
    p["depth_pref"] = str(row["Depth pref"]).strip()
    p.pop("distance_pref", None)
    hab = {}
    if num(row["Depth min"]) is not None: hab["depth_min"] = int(num(row["Depth min"], int))
    if num(row["Depth max"]) is not None: hab["depth_max"] = int(num(row["Depth max"], int))
    if num(row["Width min"]) is not None: hab["width_min"] = int(num(row["Width min"], int))
    if num(row["Width max"]) is not None: hab["width_max"] = int(num(row["Width max"], int))
    p["habitat"] = hab
    p["provinces"] = listing(row["Provinces"])
    p["biomes"] = weights(row["Biomes"])
    req = listing(row["Biomes require"])
    if req: p["biomes_require"] = req
    else: p.pop("biomes_require", None)
    p["base"] = num(row["Base"])
    if sid in HYBRIDS: p["hybrid_of"] = list(HYBRIDS[sid])
    return p


def level_model(old_profiles):
    """The old ladder stretched to 0-50, and a nearest-neighbour rule for the fish that never had a rung."""
    pts = []
    for sid, d in old_profiles.items():
        w = d.get("weight_g", {}).get("max", 1000); s = d.get("fight", {}).get("strength", 0.3)
        pts.append((math.log10(max(1, w)), s, round(d.get("min_angler_level", 0) * 50.0 / OLD_LEVEL_MAX)))
    def predict(w, s):
        x, y = math.log10(max(1, w)), s
        near = sorted(pts, key=lambda t: (t[0] - x) ** 2 + ((t[1] - y) * 2.0) ** 2)[:3]
        return int(round(sum(t[2] for t in near) / 3.0))
    return {sid: round(d.get("min_angler_level", 0) * 50.0 / OLD_LEVEL_MAX) for sid, d in old_profiles.items()}, predict


def wire_tree(tree, rows, old_main):
    name = os.path.basename(tree.rstrip("/\\"))
    prof_dir = os.path.join(tree, D, "fish_profiles")
    assets = os.path.join(tree, A)
    made = []
    stretched, predict = level_model(old_main)
    parents = {}   # parent -> {hybrid parent2: rate}
    for sid, (a, b) in HYBRIDS.items():
        parents.setdefault(a, {})[b] = 0.5
        parents.setdefault(b, {})[a] = 0.5
    ids = [r["ID"] for r in rows]
    # ---- profiles ----
    n_new = 0
    for r in rows:
        sid = r["ID"]
        pp = os.path.join(prof_dir, sid + ".json")
        old = jload(pp) if os.path.exists(pp) else None
        p = build_profile(r, old)
        for k in ("legendary", "breeds_with", "gynogenesis", "bed"):
            if old and k in old: p[k] = old[k]
        if sid in parents:
            bw = dict(p.get("breeds_with", {}))
            for other, rate in parents[sid].items(): bw.setdefault(other, rate)
            p["breeds_with"] = bw
        p["min_angler_level"] = stretched[sid] if sid in stretched else predict(p["weight_g"]["max"], p["fight"]["strength"])
        jdump(pp, p)
        if not old: n_new += 1
    made.append("profiles %d (+%d new)" % (len(rows), n_new))
    # ---- lang: names, descriptions, the new groups / diets / province ----
    for i, loc in enumerate(LOCS):
        lp = os.path.join(assets, "lang", loc + ".json"); d = jload(lp)
        ncol = ("English name", "Русское название", "Українська назва")[i]
        dcol = ("Description EN", "Описание RU", "Опис UK")[i]
        for r in rows:
            sid = r["ID"]
            nm = str(r[ncol] or r["English name"]).strip()
            d["item.riverfishing." + sid] = nm
            d["fish.riverfishing." + sid] = nm
            if r.get(dcol): d["fishdesc.riverfishing." + sid] = str(r[dcol]).strip()
        for g, names in GROUP_NAMES.items(): d["fishgroup.riverfishing." + g] = names[i]
        for k, names in DIET_NAMES.items(): d["diet.riverfishing." + k] = names[i]
        for k, names in PROVINCE_NAMES.items(): d["province.riverfishing." + k] = names[i]
        d.setdefault("journal.riverfishing.diet", ("Diet", "Питание", "Живлення")[i])
        jdump(lp, d)
    made.append("lang")
    # ---- textures (the author's art lives in the main tree) and item models ----
    tex_src = os.path.join(MAIN, A, "textures/item/fish")
    tex_dst = os.path.join(assets, "textures/item/fish")
    n_tex = 0; placeholders = []
    for sid in ids:
        src = os.path.join(tex_src, sid + ".png")
        if not os.path.exists(src):
            placeholders.append(sid)
            src = os.path.join(tex_src, "catfish.png" if "catfish" in sid else DONOR + ".png")
        dst = os.path.join(tex_dst, sid + ".png")
        if not os.path.exists(dst) or (tree != MAIN and rd_bytes(src) != rd_bytes(dst)):
            shutil.copy(src, dst); n_tex += 1
    made.append("textures +%d" % n_tex)
    n_models = 0
    for sub in ("models/item", "models/item/fish_icon", "items"):   # fish_icon: the inventory icon's own model
        donor = os.path.join(assets, sub, DONOR + ".json")
        if not os.path.exists(donor): continue
        src = rd(donor)
        for sid in ids:
            dst = os.path.join(assets, sub, sid + ".json")
            if not os.path.exists(dst):
                wr(dst, src.replace('/' + DONOR + '"', '/' + sid + '"').replace(':' + DONOR + '"', ':' + sid + '"')); n_models += 1
    made.append("models +%d" % n_models)
    # ---- roster ----
    mi = os.path.join(tree, "common/src/main/java/com/riverfishing/registry/ModItems.java"); s = rd(mi)
    missing = [x for x in ids if '"%s"' % x not in s]
    if missing:
        lines = ["            // " + TAG + ": the author's species table — %d species in one wave." % len(missing)]
        for i in range(0, len(missing), 4):
            lines.append("            " + ", ".join('"%s"' % x for x in missing[i:i + 4]) + ",")
        m = re.search(r"(public static final String\[\] FISH_SPECIES = \{\n)", s); assert m, mi
        wr(mi, s[:m.end()] + "\n".join(lines) + "\n" + s[m.end():]); made.append("roster +%d" % len(missing))
    # ---- fishes tag ----
    for cand in ("tags/item/fishes.json", "tags/items/fishes.json"):
        tp = os.path.join(tree, D, cand)
        if os.path.exists(tp):
            tag = jload(tp); have = set(tag["values"])
            add = ["riverfishing:" + x for x in ids if "riverfishing:" + x not in have]
            if add: tag["values"] += add; jdump(tp, tag); made.append("tag +%d" % len(add))
    print("  %-8s %s" % (name, ", ".join(made)))
    return placeholders


def rd_bytes(p):
    with open(p, "rb") as f: return f.read()


def main():
    args = [a for a in sys.argv[1:]]
    if not args: sys.exit(__doc__)
    table = args[0]; trees = args[1:] or TREES
    rows = read_table(table)
    prof_dir = os.path.join(MAIN, D, "fish_profiles")
    old_main = {f[:-5]: jload(os.path.join(prof_dir, f)) for f in os.listdir(prof_dir) if f.endswith(".json")}
    print("table: %d species, %d already in the main tree" % (len(rows), sum(1 for r in rows if r["ID"] in old_main)))
    for sid, (a, b) in HYBRIDS.items():
        ids = {r["ID"] for r in rows}
        assert sid in ids and a in ids and b in ids, (sid, a, b)
    ph = None
    for t in trees:
        ph = wire_tree(t, rows, old_main)
    if ph: print("placeholder textures (draw these):", ", ".join(ph))


if __name__ == "__main__":
    main()
